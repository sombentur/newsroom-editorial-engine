"""Manual Editorial Workbench (owner request, 28 Sep 2026).

The owner runs the research for Workbench topics by hand (e.g. Gemini Deep Research) and gives the app each finished
report's link through an Excel workbook:

    one Topic ID -> one research prompt -> one external research -> one report link -> one imported report

The Topic ID is the Workbench article id, so a report can never land on another topic. Downloading the workbook
reserves its topics for manual research (the scheduler does not research them itself). Uploading it saves each link
against its Topic ID and imports the reports one topic at a time: Gemini conversation links through the Chrome
extension (as Copy report again: from an open tab that shows the report, else by opening the link), other links (e.g. a
Google Doc shared with anyone who has the link) by reading the page. A topic counts as imported only when report content
was actually retrieved. The report then takes the normal route in the topic's turn (research checks, SEO, thumbnail,
publishing); nothing else is started here.
"""
import asyncio
import base64
import binascii
import io
import logging
import re
from datetime import datetime, timedelta
from html import escape
from urllib.parse import urlsplit

from lib.db import db
from lib.util import audit, new_id, now_utc

logger = logging.getLogger(__name__)

STATUS_LABELS = {
    "pending": "Pending Research",
    "excel_exported": "Excel Exported",
    "link_added": "Research Link Added",
    "ready_to_import": "Ready to Import",
    "importing": "Importing",
    "research_imported": "Research Imported",
    "link_invalid": "Link Invalid",
    "access_denied": "Access Denied",
    "report_not_found": "Report Not Found",
    "import_failed": "Import Failed",
    "research_exists": "Research Already Exists",
}
# Reserved for the owner's manual research: the scheduler leaves these topics' research to the owner.
RESERVED = {"excel_exported", "link_added", "ready_to_import", "importing", "link_invalid", "access_denied",
            "report_not_found", "import_failed"}
CLOSED_STAGES = {"rejected", "scheduled", "published", "verified", "wordpress_draft"}
MIN_REPORT_CHARS = 1500  # of readable text: less is not a research report
MAX_DOWNLOAD_BYTES = 8_000_000
COLUMNS = ["Topic ID", "Topic Title", "Website", "Language", "Category", "Research Prompt", "Research Report Link", "Status"]
GEMINI_APP = re.compile(r"^https://gemini\.google\.com/((?:u/\d{1,2}/)?)app/([A-Za-z0-9_-]{6,64})(?:[/?#].*)?$")
GEMINI_SHARE = re.compile(r"^https://(?:gemini\.google\.com/share/|g\.co/gemini/share/)")
GOOGLE_DOC = re.compile(r"^https://docs\.google\.com/document/d/([A-Za-z0-9_-]{10,})")
SHARE_MESSAGE = ("Gemini share pages load their report with scripts, so the app cannot read them. Paste the conversation "
                 "link from the address bar (https://gemini.google.com/app/…) instead.")


STOPPED_BY_EDITOR = "Stopped by editor."


def _owner_resumes(article: dict) -> dict:
    """An owner's Proceed/Fetch/upload lifts an earlier Stop (a stopped topic's browser jobs are cancelled at once)."""
    return {"held_reason": None} if article.get("held_reason") == STOPPED_BY_EDITOR else {}


class ImportProblem(Exception):
    def __init__(self, status: str, message: str):
        super().__init__(message)
        self.status = status


def has_research(article: dict) -> bool:
    return bool(article.get("dossier") or (article.get("dossier_meta") or {}).get("report") or article.get("article"))


def awaiting_manual(article: dict | None) -> bool:
    """Reserved for manual research and not imported yet: the scheduler leaves its research to the owner."""
    return bool(article and (article.get("manual_research") or {}).get("status") in RESERVED and not has_research(article))


def auto_research_enabled(system: dict | None) -> bool:
    """The Editorial Workbench switch (owner request, 28 Sep 2026): off, the app starts no research by itself."""
    return (system or {}).get("auto_research") is not False


def waits_for_link(article: dict | None, system: dict | None) -> bool:
    """Automatic research is off and this topic has no research yet: it waits for its report link."""
    return bool(article) and not auto_research_enabled(system) and not has_research(article)


async def auto_research_on() -> bool:
    return auto_research_enabled(await db.system_settings.find_one({"id": "system"}, {"auto_research": 1}))


def auto_running(article: dict) -> bool:
    """The app is researching this topic itself right now (it is not offered for manual research)."""
    return article.get("stage") == "researching" and not (article.get("manual_research") or {}).get("status")


def classify_link(link: str) -> tuple[str | None, str]:
    """(normalised URL, kind) for a report link: kind "gemini" (a Gemini conversation) or "web"; (None, reason) if
    the link cannot be used."""
    link = str(link or "").strip()
    if not link:
        return None, "empty"
    if GEMINI_SHARE.match(link):
        return None, SHARE_MESSAGE
    if match := GEMINI_APP.match(link):
        return "https://gemini.google.com/" + match.group(1) + "app/" + match.group(2), "gemini"  # its account kept
    parts = urlsplit(link)
    if parts.scheme not in {"http", "https"} or not parts.netloc or " " in link or len(link) > 2000:
        return None, "This is not a web link (it must start with https://)."
    return link.split("#", 1)[0], "web"


# ── Rows ─────────────────────────────────────────────────────────────────────
async def manual_articles() -> list[dict]:
    """Workbench topics waiting for research, plus those in the manual workflow (shown for a day once scheduled or
    published, so the owner sees them through)."""
    found = []
    recent = {"manual_research": {"$exists": True}, "stage": {"$in": ["scheduled", "published", "verified"]},
              "updated_at": {"$gte": now_utc() - timedelta(hours=24)}}
    async for article in db.articles.find({"$or": [{"stage": {"$nin": list(CLOSED_STAGES)}}, recent]}).sort("created_at", 1):
        article.pop("_id", None)
        if article.get("manual_research") or not has_research(article):
            found.append(article)
    return found


async def _prompt_builder():
    from lib.report_article import DEEP_RESEARCH_ARTICLE_PROMPT, research_prompt
    from lib.workflow import _get_prompt
    template = await _get_prompt("deep_research_article", DEEP_RESEARCH_ARTICLE_PROMPT)
    sites = {s["key"]: s async for s in db.sites.find({})}
    stamp = now_utc().isoformat()

    def prompt(article: dict) -> str:
        saved = (article.get("manual_research") or {}).get("prompt") or (article.get("dossier_meta") or {}).get("prompt")
        site = sites.get(article.get("site_key")) or {}
        if saved or not site:
            return saved or ""
        return research_prompt(template, site, article.get("topic_snapshot") or {"topic": ""}, stamp)
    return prompt, sites


def _status(article: dict) -> str:
    return (article.get("manual_research") or {}).get("status") or "pending"


def row_view(article: dict, site: dict, prompt: str, busy: bool = False) -> dict:
    """busy: a browser job of the app's own (research, SEO or thumbnail) is open for this topic right now."""
    record = article.get("manual_research") or {}
    status = _status(article)
    topic = article.get("topic_snapshot") or {}
    return {
        "id": article["id"], "title": ((article.get("article") or {}).get("headline") or topic.get("topic") or "")[:300],
        "site_key": article.get("site_key"), "website": site.get("name") or article.get("site_key"),
        "language": "Kannada" if site.get("language") == "kn" else "English", "category": topic.get("category") or "",
        "prompt": prompt, "link": record.get("link") or "", "pending_link": record.get("pending_link") or "",
        "status": status, "status_label": STATUS_LABELS.get(status, status), "message": record.get("message") or "",
        "imported_at": record.get("imported_at"), "chars": record.get("chars"), "source": record.get("source"),
        "stage": article.get("stage"), "held_reason": article.get("held_reason"), "proceed_at": record.get("proceed_at"),
        "closed": article.get("stage") in CLOSED_STAGES,
        "has_research": has_research(article), "auto_running": auto_running(article) or busy,
        "exportable": not has_research(article) and not (auto_running(article) or busy) and status != "importing",
    }


async def _busy_topics() -> set[str]:
    """Topics with an open browser job of the app's own: never offered for manual research while it works."""
    return {job["article_id"] async for job in db.browser_jobs.find(
        {"status": {"$in": ["queued", "running", "attention"]}, "manual_import": {"$ne": True}}, {"article_id": 1})}


async def row_for(article_id: str) -> dict | None:
    """One post's research row (the Editorial Workbench shows the same box as the Manual Workbench)."""
    article = await db.articles.find_one({"id": article_id})
    if not article:
        return None
    article.pop("_id", None)
    prompt, sites = await _prompt_builder()
    return row_view(article, sites.get(article.get("site_key")) or {}, prompt(article), article_id in await _busy_topics())


async def list_rows() -> list[dict]:
    prompt, sites = await _prompt_builder()
    busy = await _busy_topics()
    return [row_view(a, sites.get(a.get("site_key")) or {}, prompt(a), a["id"] in busy) for a in await manual_articles()]


# ── Excel export ─────────────────────────────────────────────────────────────
def build_workbook(rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    book = Workbook()
    sheet = book.active
    sheet.title = "Research"
    sheet.append(COLUMNS)
    for row in rows:
        sheet.append([row["id"], row["title"], row["website"], row["language"], row["category"], row["prompt"],
                      row["link"], row["status_label"]])
    header_fill = PatternFill("solid", fgColor="1F2A44")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    link_fill = PatternFill("solid", fgColor="FFF7E0")  # the column the owner fills in
    for index, width in enumerate([38, 48, 24, 11, 18, 90, 55, 20], start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (2, 6))
        cells[6].fill = link_fill
        sheet.row_dimensions[cells[0].row].height = 120
    sheet.freeze_panes = "B2"
    guide = book.create_sheet("How to use")
    for line in [
        "Manual Editorial Workbench — research workbook",
        "1. For each row, copy the Research Prompt and run it on its own (e.g. a new Gemini Deep Research chat).",
        "2. When the report is finished, copy its link: the Gemini conversation address (https://gemini.google.com/app/…),",
        "   or a Google Doc / web page shared with anyone who has the link.",
        "3. Paste the link into the Research Report Link cell of the SAME row. Keep the Topic ID unchanged.",
        "4. Save the file and upload it in the Manual Editorial Workbench (Upload Research Excel).",
        "Rows without a link simply stay Pending Research. One report link belongs to one Topic ID only.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    guide["A1"].font = Font(bold=True, size=13)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


async def export_workbook(actor: str) -> tuple[bytes, int]:
    """The workbook of every topic available for manual research; those topics are reserved for manual research."""
    prompt, sites = await _prompt_builder()
    busy = await _busy_topics()
    rows = []
    for article in await manual_articles():
        view = row_view(article, sites.get(article.get("site_key")) or {}, prompt(article), article["id"] in busy)
        if not view["exportable"] or not view["prompt"]:
            continue
        record = dict(article.get("manual_research") or {})
        record["prompt"] = view["prompt"]  # the prompt in the workbook stays the topic's prompt
        if _status(article) == "pending":
            record.update(status="excel_exported", exported_at=now_utc(), message="In the research workbook")
            view.update(status="excel_exported", status_label=STATUS_LABELS["excel_exported"])
        await db.articles.update_one({"id": article["id"]}, {"$set": {"manual_research": record}})
        rows.append(view)
    await audit("manual_research_exported", "article", None, None, actor=actor, detail={"topics": len(rows)})
    return build_workbook(rows), len(rows)


# ── Excel import ─────────────────────────────────────────────────────────────
def read_workbook(content: bytes) -> tuple[list[dict], list[str]]:
    """Rows of the uploaded workbook: {row, topic_id, link, title}. Errors when it cannot be used at all."""
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        return [], ["This file is not an Excel workbook (.xlsx). Upload the file downloaded from this page."]
    sheet = book["Research"] if "Research" in book.sheetnames else book.worksheets[0]
    lines = sheet.iter_rows()
    header = next(lines, None) or []
    names = {re.sub(r"\s+", " ", str(cell.value or "")).strip().lower(): index for index, cell in enumerate(header)}
    errors = [f"The workbook has no “{label}” column." for label in ("Topic ID", "Research Report Link")
              if label.lower() not in names]
    if errors:
        return [], errors
    id_col, link_col, title_col = names["topic id"], names["research report link"], names.get("topic title")
    rows = []
    for cells in lines:
        def value(index):
            return cells[index].value if index is not None and index < len(cells) else None
        topic_id = str(value(id_col) or "").strip()
        link = str(value(link_col) or "").strip()
        cell = cells[link_col] if link_col < len(cells) else None
        if cell is not None and cell.hyperlink and cell.hyperlink.target and not link.lower().startswith("http"):
            link = str(cell.hyperlink.target).strip()  # a link inserted with its own display text
        if not topic_id and not link:
            continue
        rows.append({"row": cells[0].row, "topic_id": topic_id, "link": link, "title": str(value(title_col) or "")[:300]})
    return rows, []


async def import_workbook(content: bytes, filename: str, actor: str) -> dict:
    """Save every valid row's link against its Topic ID; start importing. Errors in one row never stop the others."""
    rows, errors = read_workbook(content)
    batch = {"id": new_id(), "filename": str(filename or "")[:120], "created_at": now_utc(), "actor": actor,
             "errors": errors, "rows": []}
    ids = [r["topic_id"] for r in rows]
    duplicate_ids = {t for t in ids if t and ids.count(t) > 1}
    normalised = [classify_link(r["link"])[0] for r in rows]
    duplicate_links = {u for u in normalised if u and normalised.count(u) > 1}
    for entry, url in zip(rows, normalised):
        result = {"row": entry["row"], "topic_id": entry["topic_id"], "title": entry["title"], "link": entry["link"],
                  "link_state": "Added" if entry["link"] else "Empty", "matched": False, "queued": False}
        batch["rows"].append(result)
        topic_id = entry["topic_id"]
        if not topic_id:
            result.update(result="Missing Topic ID", error=True)
            continue
        if topic_id in duplicate_ids:
            result.update(result="Duplicate Topic ID (row skipped)", error=True)
            continue
        article = await db.articles.find_one({"id": topic_id})
        if not article:
            result.update(result="Unknown Topic ID (not created)", error=True)
            continue
        result.update(matched=True, title=result["title"] or ((article.get("topic_snapshot") or {}).get("topic") or "")[:300])
        if article.get("stage") in CLOSED_STAGES:
            result.update(result="Topic is already published or rejected (skipped)", error=True)
            continue
        record = dict(article.get("manual_research") or {})
        if not entry["link"]:
            result["result"] = STATUS_LABELS.get(record.get("status") or "pending", "Pending Research")
            continue
        kind = classify_link(entry["link"])[1]
        if not url:
            record.update(status="link_invalid", link=entry["link"][:2000], message=kind)
            await db.articles.update_one({"id": topic_id}, {"$set": {"manual_research": record}})
            result.update(result=STATUS_LABELS["link_invalid"], message=kind, error=True)
            continue
        if url in duplicate_links:
            result.update(result="The same link is in more than one row (row skipped)", error=True)
            continue
        other = await db.articles.find_one({"manual_research.link": url, "id": {"$ne": topic_id}}, {"id": 1})
        if other:
            result.update(result=f"This link already belongs to Topic {other['id']} (row skipped)", error=True)
            continue
        if has_research(article):
            if record.get("link") == url and record.get("status") == "research_imported":
                result["result"] = STATUS_LABELS["research_imported"]
                continue
            record.update(status="research_exists", pending_link=url,
                          message="This topic already has research. Use Replace Existing Research to import this link.")
            await db.articles.update_one({"id": topic_id}, {"$set": {"manual_research": record}})
            result.update(result=STATUS_LABELS["research_exists"])
            continue
        if record.get("status") == "importing" and record.get("link") == url:
            result.update(result=STATUS_LABELS["importing"], queued=True)
            continue
        record.update(status="ready_to_import", link=url, kind=kind, queued_at=now_utc(), message="Waiting to be fetched")
        record.pop("pending_link", None)
        await db.articles.update_one({"id": topic_id}, {"$set": {"manual_research": record, **_owner_resumes(article)}})
        result.update(result=STATUS_LABELS["ready_to_import"], queued=True)
    await db.manual_imports.insert_one(dict(batch))
    await audit("manual_research_uploaded", "article", None, None, actor=actor,
                detail={"rows": len(rows), "queued": sum(r["queued"] for r in batch["rows"]), "errors": errors})
    ensure_runner()
    return await batch_summary(batch["id"])


def decode_upload(content_b64: str) -> bytes:
    try:
        content = base64.b64decode(content_b64, validate=True)
    except (binascii.Error, ValueError):
        raise ValueError("The uploaded file could not be read.") from None
    if not content or len(content) > 10_000_000:
        raise ValueError("Upload an Excel workbook below 10 MB.")
    return content


async def batch_summary(batch_id: str) -> dict | None:
    """The upload's summary with each row's current import result (imports run after the upload)."""
    batch = await db.manual_imports.find_one({"id": batch_id}, {"_id": 0})
    if not batch:
        return None
    counts = {"total_rows": len(batch["rows"]), "matched": 0, "links_found": 0, "imported": 0, "invalid_links": 0,
              "access_denied": 0, "not_found": 0, "failed": 0, "without_links": 0, "row_errors": 0, "in_progress": 0,
              "already_exists": 0}
    for row in batch["rows"]:
        if row.get("queued"):
            article = await db.articles.find_one({"id": row["topic_id"]}, {"manual_research": 1}) or {}
            record = article.get("manual_research") or {}
            status = record.get("status") or "pending"
            row.update(result=STATUS_LABELS.get(status, status), message=record.get("message") or "", status=status)
        status = row.get("status") or ""
        result = row.get("result") or ""
        counts["matched"] += bool(row.get("matched"))
        counts["links_found"] += bool(row["link"])
        counts["without_links"] += not row["link"] and bool(row.get("matched"))
        counts["imported"] += status == "research_imported"
        counts["in_progress"] += status in {"ready_to_import", "importing"}
        counts["invalid_links"] += result == STATUS_LABELS["link_invalid"] or status == "link_invalid"
        counts["access_denied"] += status == "access_denied"
        counts["not_found"] += status == "report_not_found"
        counts["failed"] += status == "import_failed"
        counts["already_exists"] += result == STATUS_LABELS["research_exists"]
        counts["row_errors"] += bool(row.get("error")) and result != STATUS_LABELS["link_invalid"]
    batch["counts"] = counts
    return batch


# ── Fetching the reports ─────────────────────────────────────────────────────
_runner: asyncio.Task | None = None


def ensure_runner() -> None:
    """Import the queued reports, one topic at a time (a single background task)."""
    global _runner
    if _runner is None or _runner.done():
        _runner = asyncio.get_running_loop().create_task(_run())


async def _run() -> None:
    # An import interrupted by a restart is queued again (no other import runs while this task starts).
    await db.articles.update_many({"manual_research.status": "importing"},
                                  {"$set": {"manual_research.status": "ready_to_import",
                                            "manual_research.message": "Interrupted; fetching again"}})
    while True:
        article = await db.articles.find_one({"manual_research.status": "ready_to_import"},
                                             sort=[("manual_research.queued_at", 1)])
        if not article:
            return
        try:
            await import_one(article)
        except Exception:
            logger.exception("manual research import failed for %s", article.get("id"))
            await _set(article["id"], "import_failed", "The import stopped unexpectedly. Use Fetch Research to try again.")


async def _set(article_id: str, status: str, message: str, **fields) -> None:
    update = {"manual_research.status": status, "manual_research.message": message[:400],
              **{f"manual_research.{k}": v for k, v in fields.items()}}
    await db.articles.update_one({"id": article_id}, {"$set": update})


async def import_one(article: dict) -> None:
    record = article.get("manual_research") or {}
    url = record.get("link") or ""
    await _set(article["id"], "importing", "Fetching the report…", started_at=now_utc())
    try:
        if GEMINI_APP.match(url):
            content, source = await _from_gemini(article, url), "Gemini conversation"
        else:
            content, source = await _from_web(url), "web link"
    except ImportProblem as problem:
        await _set(article["id"], problem.status, str(problem))
        return
    text = plain_text(content)
    if len(text) < MIN_REPORT_CHARS:
        await _set(article["id"], "report_not_found",
                   f"The link opened, but it held only {len(text)} characters of text: no research report was found there.")
        return
    current = await db.articles.find_one({"id": article["id"]}) or {}
    if (current.get("manual_research") or {}).get("status") != "importing" or current.get("stage") in CLOSED_STAGES:
        return  # released, rejected or published meanwhile: nothing is overwritten
    await store_report(current, url, content, source)


def plain_text(content: str) -> str:
    return " ".join(re.sub(r"<[^>]+>", " ", re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", content or "")).split())


async def store_report(article: dict, url: str, content: str, source: str) -> None:
    """The retrieved report becomes this topic's research (the same record a Deep Research copy fills)."""
    record = dict(article.get("manual_research") or {})
    prompt = record.get("prompt") or (article.get("dossier_meta") or {}).get("prompt") or ""
    chars = len(plain_text(content))
    record.update(status="research_imported", link=url, source=source, imported_at=now_utc(), chars=chars,
                  message=f"Imported {chars:,} characters from the {source}")
    record.pop("pending_link", None)
    update = {"manual_research": record, "dossier": None, "validation": None, "updated_at": now_utc(),
              "dossier_meta": {"report_mode": True, "prompt": prompt, "report": content, "provider": "manual-research",
                               "model": source, "status": "formatting", "started": now_utc(), "manual_link": url}}
    if article.get("stage") in {"held_review", "failed"}:
        update.update(stage="selected", held_reason=None, ai_failure=None)
    await db.articles.update_one({"id": article["id"]}, {"$set": update, "$push": {"history": {
        "stage": update.get("stage", article.get("stage")), "at": now_utc(), "actor": "owner",
        "note": f"Research imported from the report link ({source}): {url}. It continues with the research checks, "
                "SEO and thumbnail in its turn."}}})
    await audit("manual_research_imported", "article", article["id"], article.get("site_key"),
                detail={"source": source, "chars": chars})



async def today_stats() -> dict:
    """Count of articles that reached published/scheduled/wordpress_draft today (IST) per site."""
    from datetime import timedelta, timezone as _tz
    ist = _tz(timedelta(hours=5, minutes=30))
    start_ist = now_utc().astimezone(ist).replace(hour=0, minute=0, second=0, microsecond=0)
    start_utc = start_ist.astimezone(_tz.utc)
    pipeline = [
        {"$match": {"stage": {"$in": ["published", "scheduled", "wordpress_draft"]},
                    "updated_at": {"$gte": start_utc}}},
        {"$group": {"_id": "$site_key", "count": {"$sum": 1}}},
    ]
    docs = await db.articles.aggregate(pipeline).to_list(None)
    return {(d["_id"] or "unknown"): d["count"] for d in docs}



async def import_pasted(article_id: str, html: str, replace: bool = False) -> dict:
    """Store HTML content the owner pasted directly as this topic's research."""
    import re as _re
    article = await db.articles.find_one({"id": article_id})
    if not article:
        raise LookupError("Topic not found.")
    if article.get("stage") in CLOSED_STAGES:
        raise ValueError("This topic is already scheduled, published or rejected.")
    if has_research(article) and not replace:
        raise ValueError("This topic already has research. Confirm replacing it.")
    cleaned = main_html(html) if _re.match(r"\s*<", html) else f"<p>{escape(html)}</p>"
    return await _store_direct(article, cleaned, "pasted content", "pasted-content")


async def import_uploaded(article_id: str, filename: str, content_b64: str, replace: bool = False) -> dict:
    """Parse a .docx or .pdf file and store its text as this topic's research."""
    article = await db.articles.find_one({"id": article_id})
    if not article:
        raise LookupError("Topic not found.")
    if article.get("stage") in CLOSED_STAGES:
        raise ValueError("This topic is already scheduled, published or rejected.")
    if has_research(article) and not replace:
        raise ValueError("This topic already has research. Confirm replacing it.")
    try:
        raw = base64.b64decode(content_b64)
    except Exception:
        raise ValueError("Could not decode the uploaded file.") from None
    lname = (filename or "").lower()
    if lname.endswith(".docx"):
        html = _docx_to_html(raw)
        source = f"Word document ({filename})"
    elif lname.endswith(".pdf"):
        html = _pdf_to_html(raw)
        source = f"PDF document ({filename})"
    else:
        raise ValueError("Only Word (.docx) and PDF (.pdf) files are accepted.")
    return await _store_direct(article, html, source, f"uploaded:{filename}")


async def _store_direct(article: dict, html: str, source: str, virtual_url: str) -> dict:
    """Validate then store directly provided HTML as the article's research."""
    text = plain_text(html)
    if len(text) < MIN_REPORT_CHARS:
        raise ValueError(
            f"The content has only {len(text):,} readable characters; "
            f"a research report needs at least {MIN_REPORT_CHARS:,}. "
            "Paste or upload the full research report."
        )
    record = dict(article.get("manual_research") or {})
    prompt = record.get("prompt") or (article.get("dossier_meta") or {}).get("prompt") or ""
    chars = len(text)
    record.update(status="research_imported", link=virtual_url, source=source,
                  imported_at=now_utc(), chars=chars,
                  message=f"Imported {chars:,} characters from {source}")
    record.pop("pending_link", None)
    update = {
        "manual_research": record,
        "dossier": None, "validation": None, "article": None,
        "article_meta": None, "quality_gate": None, "image": None,
        "updated_at": now_utc(),
        "dossier_meta": {
            "report_mode": True, "prompt": prompt, "report": html,
            "provider": "manual-research", "model": source,
            "status": "formatting", "started": now_utc(), "manual_link": virtual_url,
        },
        **_owner_resumes(article),
    }
    if article.get("stage") in {"held_review", "failed"}:
        update.update(stage="selected", held_reason=None, ai_failure=None)
    await db.articles.update_one(
        {"id": article["id"]},
        {"$set": update, "$push": {"history": {
            "stage": update.get("stage", article.get("stage")),
            "at": now_utc(), "actor": "owner",
            "note": (f"Research imported directly from {source}. "
                     "Continues with research checks, SEO and thumbnail in its turn."),
        }}},
    )
    await audit("manual_research_imported", "article", article["id"], article.get("site_key"),
                detail={"source": source, "chars": chars})
    return {"ok": True, "chars": chars}


def _docx_to_html(content: bytes) -> str:
    """Convert a .docx file to HTML preserving paragraphs, headings and tables."""
    from docx import Document
    from io import BytesIO

    doc = Document(BytesIO(content))

    def _runs(para) -> str:
        parts = []
        for run in para.runs:
            t = escape(run.text or "")
            if run.bold:
                t = f"<b>{t}</b>"
            elif run.italic:
                t = f"<i>{t}</i>"
            parts.append(t)
        return "".join(parts).strip()

    def _para(para) -> str:
        text = _runs(para)
        if not text:
            return ""
        style = (para.style.name or "").lower()
        for lvl in range(1, 5):
            if f"heading {lvl}" in style:
                return f"<h{lvl}>{text}</h{lvl}>"
        return f"<p>{text}</p>"

    def _table(table) -> str:
        rows = []
        for i, row in enumerate(table.rows):
            tag = "th" if i == 0 else "td"
            cells = "".join(
                f"<{tag}>{escape(c.text.strip())}</{tag}>" for c in row.cells
            )
            rows.append(f"<tr>{cells}</tr>")
        return f'<table border="1">{"".join(rows)}</table>'

    tbl_map = {t._element: t for t in doc.tables}
    para_map = {p._element: p for p in doc.paragraphs}
    parts = []
    for child in doc.element.body:
        tag_name = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag_name == "tbl" and child in tbl_map:
            parts.append(_table(tbl_map[child]))
        elif tag_name == "p" and child in para_map:
            h = _para(para_map[child])
            if h:
                parts.append(h)
    return "\n".join(parts) if parts else "<p>(Empty document)</p>"


def _pdf_to_html(content: bytes) -> str:
    """Extract text from a PDF (table structure not preserved)."""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ValueError(
            "PDF support requires the pypdf package. Use a Word .docx file instead."
        ) from None
    from io import BytesIO

    reader = PdfReader(BytesIO(content))
    parts = []
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        for para in text.split("\n\n"):
            para = para.strip()
            if para:
                parts.append(f"<p>{escape(para)}</p>")
    return "\n".join(parts) if parts else "<p>(No text content found in PDF)</p>"


async def _from_gemini(article: dict, url: str) -> str:
    """A Gemini conversation's finished report, through the Chrome extension (as Copy report again)."""
    from lib.ai import AIError
    from lib.browser_bridge import _awaiting, _wait_for_job
    config = await db.browser_config.find_one({"id": "browser"}) or {}
    tab_id = ((config.get("workspace") or {}).get("research") or {}).get("tab_id")
    if not tab_id:
        raise ImportProblem("import_failed", "Connect the Gemini work tab first (Setup Wizard → Browser Extension).")
    record = article.get("manual_research") or {}
    job = {"id": new_id(), "article_id": article["id"], "kind": "research",
           "prompt": record.get("prompt") or (article.get("topic_snapshot") or {}).get("topic") or "manual research",
           "status": "queued", "message": "Importing the manual research report from its Gemini conversation",
           "capture_existing": True, "manual_import": True, "capture_tab_id": tab_id, "capture_url": url,
           "capture_candidates": [url], "created_at": now_utc()}
    await db.browser_jobs.insert_one(dict(job))
    _awaiting[job["id"]] = _awaiting.get(job["id"], 0) + 1
    try:
        result = await _wait_for_job(job, "research", 0)
    except AIError as exc:
        # A job that stopped must not keep the Gemini lane.
        await db.browser_jobs.update_one({"id": job["id"], "status": {"$in": ["queued", "running", "attention"]}},
                                         {"$set": {"status": "cancelled", "message": "Manual research import ended"}})
        message = str(exc)
        if message == job["message"]:
            message = ("The fetch was stopped before the report could be copied (the topic was stopped, or the app was "
                       "paused). Press Proceed or Fetch Research to try again.")
        if re.search(r"no finished report|not progressing", message):
            raise ImportProblem("report_not_found", message) from None
        raise ImportProblem("import_failed", message) from None
    finally:
        _awaiting[job["id"]] -= 1
        if _awaiting[job["id"]] <= 0:
            _awaiting.pop(job["id"], None)
    await db.browser_jobs.update_one({"id": job["id"]}, {"$set": {"status": "consumed"}, "$unset": {"result": ""}})
    return result


async def _from_web(url: str, client=None) -> str:
    """The report at a public link (a Google Doc is read through its HTML export)."""
    import httpx
    fetch = url
    if match := GOOGLE_DOC.match(url):
        fetch = f"https://docs.google.com/document/d/{match.group(1)}/export?format=html"
    try:
        if client is None:
            async with httpx.AsyncClient(follow_redirects=True, timeout=30,
                                         headers={"User-Agent": "Mozilla/5.0 (Newsroom research import)"}) as own:
                response = await own.get(fetch)
        else:
            response = await client.get(fetch)
    except httpx.HTTPError as exc:
        raise ImportProblem("import_failed", f"The link could not be opened ({type(exc).__name__}).") from None
    final = str(response.url)
    if response.status_code in (401, 403) or urlsplit(final).netloc == "accounts.google.com" or "ServiceLogin" in final:
        raise ImportProblem("access_denied", "The report is not public. Share it with anyone who has the link, or use "
                                             "its Gemini conversation link.")
    if response.status_code in (404, 410):
        raise ImportProblem("report_not_found", f"The link returned {response.status_code}: there is no page there.")
    if response.status_code >= 400:
        raise ImportProblem("import_failed", f"The link returned HTTP {response.status_code}.")
    kind = response.headers.get("content-type", "")
    if len(response.content) > MAX_DOWNLOAD_BYTES:
        raise ImportProblem("import_failed", "The page is too large to import.")
    if "html" in kind:
        return main_html(response.text)
    if kind.startswith("text/"):
        return response.text
    raise ImportProblem("import_failed", f"The link is not a web page or document ({kind or 'unknown type'}).")


def main_html(html: str) -> str:
    """The page's main content (its article or main section), without scripts, menus or page furniture."""
    from lxml import html as lxml_html
    try:
        root = lxml_html.fromstring(html)
    except Exception:
        return "<p>" + escape(plain_text(html)) + "</p>"
    for node in root.xpath("//script|//style|//noscript|//nav|//header|//footer|//aside|//form|//svg|//iframe"):
        node.drop_tree()
    main = next(iter(root.xpath("//article") or root.xpath("//main") or root.xpath("//*[@role='main']")
                     or root.xpath("//body") or [root]))
    return lxml_html.tostring(main, encoding="unicode")


async def fetch_again(article_id: str, replace: bool) -> dict:
    """Fetch / Re-import Research, or Replace Existing Research (explicit owner actions)."""
    article = await db.articles.find_one({"id": article_id})
    if not article:
        raise LookupError("Topic not found.")
    if article.get("stage") in CLOSED_STAGES:
        raise ValueError("This topic is already published or rejected.")
    record = dict(article.get("manual_research") or {})
    url = record.get("pending_link") if record.get("status") == "research_exists" else record.get("link")
    if not url or not classify_link(url)[0]:
        raise ValueError("This topic has no usable research report link yet.")
    if record.get("status") == "importing":
        raise ValueError("This report is being fetched right now.")
    if has_research(article) and not replace:
        raise ValueError("This topic already has research. Use Replace Existing Research (or Re-import) to overwrite it.")
    reset = {}
    if has_research(article):
        # The new report replaces the research, so the article, SEO and thumbnail made from the old one are made again.
        reset = {"dossier": None, "validation": None, "article": None, "article_meta": None, "quality_gate": None,
                 "image": None, "stage": "selected", "held_reason": None, "ai_failure": None,
                 "dossier_meta": {**(article.get("dossier_meta") or {}), "report": None}}
    record.update(status="ready_to_import", link=url, queued_at=now_utc(), message="Waiting to be fetched")
    record.pop("pending_link", None)
    note = ("Replacing the existing research with the report at " if reset else "Fetching the research report from ") + url
    await db.articles.update_one({"id": article_id}, {
        "$set": {"manual_research": record, **_owner_resumes(article), **reset, "updated_at": now_utc()},
        "$push": {"history": {"stage": reset.get("stage", article.get("stage")), "at": now_utc(), "actor": "owner",
                              "note": note}}})
    ensure_runner()
    return record


def proceed_key(article: dict) -> tuple:
    """Sort key: topics the owner pressed Proceed on first (earliest first); the rest keep their order."""
    pressed = (article.get("manual_research") or {}).get("proceed_at")
    return (0, pressed.timestamp()) if isinstance(pressed, datetime) else (1, 0.0)


async def proceed(article_id: str, url: str | None = None) -> dict:
    """Proceed (Manual Workbench): the report at the pasted (or saved) link becomes this topic's research, and the
    topic goes next: SEO, thumbnail and publishing/scheduling follow in the normal way."""
    article = await db.articles.find_one({"id": article_id})
    if not article:
        raise LookupError("Topic not found.")
    if article.get("stage") in CLOSED_STAGES:
        raise ValueError("This topic is already scheduled, published or rejected.")
    record = dict(article.get("manual_research") or {})
    status = record.get("status")
    note = "Proceed: this topic goes next (SEO, thumbnail, then publishing or scheduling)"
    if url and url.strip():
        link, kind = classify_link(url)
        if not link:
            raise ValueError(kind if kind != "empty" else "Paste the research report link first.")
        other = await db.articles.find_one({"manual_research.link": link, "id": {"$ne": article_id}},
                                           {"id": 1, "topic_snapshot.topic": 1})
        if other:
            raise ValueError("This link already belongs to another topic: “"
                             + ((other.get("topic_snapshot") or {}).get("topic") or other["id"])[:80] + "”.")
        if has_research(article):
            if not (record.get("link") == link and status == "research_imported"):
                raise ValueError("This topic already has research. Use Replace Existing Research to use another report.")
        elif status == "importing":
            raise ValueError("A report is being fetched for this topic right now.")
        else:
            record.update(status="ready_to_import", link=link, kind=kind, queued_at=now_utc(),
                          message="Waiting to be fetched; the topic goes next once its report is in")
            record.pop("pending_link", None)
            note = (f"Proceed with the research report at {link}: it is fetched, then this topic goes next (SEO, "
                    "thumbnail, then publishing or scheduling)")
    elif not has_research(article):
        if not record.get("link") or not classify_link(record["link"])[0]:
            raise ValueError("Paste the research report link first.")
        if status not in {"ready_to_import", "importing"}:
            record.update(status="ready_to_import", queued_at=now_utc(),
                          message="Waiting to be fetched; the topic goes next once its report is in")
    record["proceed_at"] = now_utc()
    await db.articles.update_one({"id": article_id}, {
        "$set": {"manual_research": record, **_owner_resumes(article), "updated_at": now_utc()},
        "$push": {"history": {"stage": article.get("stage"), "at": now_utc(), "actor": "owner", "note": note}}})
    ensure_runner()
    return record


async def import_again(article_id: str, url: str) -> dict:
    """Copy report again / the Editorial Workbench link box (owner report, 29 Sep 2026): the report at this Gemini link
    becomes the topic's research now, whatever holds the queue, and the topic goes next. Research it had is replaced
    (the owner asked for this copy): its article, SEO and thumbnail are made again, and its other jobs give way."""
    article = await db.articles.find_one({"id": article_id})
    if not article:
        raise LookupError("Article not found.")
    if article.get("stage") in CLOSED_STAGES or (article.get("wp") or {}).get("post_id"):
        raise ValueError("Published or scheduled articles are locked.")
    link, kind = classify_link(url)
    if not link:
        raise ValueError(kind)
    other = await db.articles.find_one({"manual_research.link": link, "id": {"$ne": article_id}},
                                       {"id": 1, "topic_snapshot.topic": 1})
    if other:
        raise ValueError("This link already belongs to another topic: “"
                         + ((other.get("topic_snapshot") or {}).get("topic") or other["id"])[:80] + "”.")
    record = dict(article.get("manual_research") or {})
    if record.get("status") == "importing":
        raise ValueError("A report is being fetched for this topic right now.")
    now = now_utc()
    record.update(status="ready_to_import", link=link, kind=kind, queued_at=now, proceed_at=now,
                  message="Waiting to be fetched; the topic goes next once its report is in")
    record.pop("pending_link", None)
    meta = {**(article.get("dossier_meta") or {}), "report": None}
    meta.pop("recopy_url", None)
    await db.browser_jobs.update_many(
        {"article_id": article_id, "status": {"$in": ["queued", "running", "attention"]}, "manual_import": {"$ne": True}},
        {"$set": {"status": "cancelled", "message": "Replaced by a report re-copy"}})
    await db.articles.update_one({"id": article_id}, {"$set": {
        "manual_research": record, "dossier_meta": meta, "dossier": None, "validation": None, "article": None,
        "article_meta": None, "quality_gate": None, "image": None, "stage": "selected", "held_reason": None,
        "ai_failure": None, "updated_at": now},
        "$push": {"history": {"stage": "selected", "at": now, "actor": "owner",
                              "note": f"Copying the finished research report from {link} now; the topic then goes next"}}})
    ensure_runner()
    return record


async def release(article_id: str) -> None:
    """Back to automatic research: the topic leaves the manual workflow (never while importing or once imported)."""
    article = await db.articles.find_one({"id": article_id})
    if not article:
        raise LookupError("Topic not found.")
    if _status(article) in {"importing", "research_imported"} or has_research(article):
        raise ValueError("This topic already has (or is importing) its research.")
    await db.articles.update_one({"id": article_id}, {"$unset": {"manual_research": ""}, "$push": {"history": {
        "stage": article.get("stage"), "at": now_utc(), "actor": "owner",
        "note": "Returned to automatic research (removed from the manual research workbook)"}}})


async def imported_report(article_id: str) -> dict | None:
    article = await db.articles.find_one({"id": article_id}, {"_id": 0, "id": 1, "manual_research": 1,
                                                             "dossier_meta": 1, "topic_snapshot.topic": 1})
    if not article:
        return None
    meta = article.get("dossier_meta") or {}
    return {"id": article["id"], "title": (article.get("topic_snapshot") or {}).get("topic"),
            "link": (article.get("manual_research") or {}).get("link"), "report": meta.get("report") or "",
            "source": meta.get("model")}
