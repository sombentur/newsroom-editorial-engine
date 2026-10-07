"""Create a polished, portable Word research dossier from a validated article record."""

from io import BytesIO
import re
from urllib.parse import urlparse

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "17365D"
BLUE = "DCE6F1"
PALE = "F5F7FA"
GREY = RGBColor(89, 89, 89)


def _shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _table(doc, headers: list[str], rows: list[list[str]]):
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for i, heading in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = heading
        _shade(cell, NAVY)
        _set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    _repeat_header(table.rows[0])
    for row_no, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value or "—")
            _set_cell_margins(cells[i])
            if row_no % 2:
                _shade(cells[i], PALE)
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(8.5)
    doc.add_paragraph()


def _bullets(doc, values):
    for value in values or []:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(str(value))


def _section(doc, title: str, text: str | None = None):
    doc.add_heading(title, level=1)
    if text:
        doc.add_paragraph(str(text))


def _plain_markdown(text: str) -> str:
    """Remove lightweight Markdown markers while preserving the report's wording."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*([^*]+?)\*(?!\*)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


def _pipe_cells(line: str) -> list[str]:
    return [_plain_markdown(cell.strip()) for cell in line.strip().strip("|").split("|")]


def _is_table_separator(line: str) -> bool:
    cells = _pipe_cells(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def _add_original_report(doc, report: str):
    _section(doc, "Original Deep Research Report")
    lines = (report or "").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        # Gemini commonly returns Markdown pipe tables. Turn each complete block into
        # an actual Word table so columns stay aligned and repeat across page breaks.
        if line.startswith("|") and i + 1 < len(lines) and _is_table_separator(lines[i + 1].strip()):
            headers = _pipe_cells(line)
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                values = _pipe_cells(lines[i])
                if len(values) < len(headers):
                    values += [""] * (len(headers) - len(values))
                rows.append(values[:len(headers)])
                i += 1
            _table(doc, headers, rows)
            continue
        if line.startswith("### "):
            doc.add_heading(_plain_markdown(line[4:]), level=3)
        elif line.startswith("## "):
            doc.add_heading(_plain_markdown(line[3:]), level=2)
        elif line.startswith("# "):
            doc.add_heading(_plain_markdown(line[2:]), level=1)
        elif line.startswith(("- ", "* ")):
            doc.add_paragraph(_plain_markdown(line[2:]), style="List Bullet")
        else:
            p = doc.add_paragraph(_plain_markdown(line))
            if line.startswith("*") and line.endswith("*"):
                for run in p.runs:
                    run.italic = True
        i += 1


def build_research_docx(article: dict, site: dict) -> bytes:
    dossier = article.get("dossier") or {}
    meta = article.get("dossier_meta") or {}
    topic = article.get("topic_snapshot") or {}
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Nirmala UI" if site.get("language") == "kn" else "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.name = normal.font.name
        styles[name].font.color.rgb = RGBColor(23, 54, 93)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(str(topic.get("topic") or "Research Dossier"))
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"Research dossier · {site.get('name', article.get('site_key', ''))}")
    run.bold = True
    run.font.color.rgb = GREY
    details = doc.add_paragraph()
    details.alignment = WD_ALIGN_PARAGRAPH.CENTER
    details.add_run(f"Model: {meta.get('model', '—')}  |  Status: {meta.get('status', 'validated')}").italic = True
    doc.add_paragraph()

    _section(doc, "Executive Summary", dossier.get("executive_summary"))
    _section(doc, "Verified Facts")
    _bullets(doc, dossier.get("verified_facts"))
    if dossier.get("newest_development"):
        _section(doc, "Latest Development", dossier.get("newest_development"))
    _section(doc, "Timeline")
    _table(doc, ["Date", "Event"], [[x.get("date"), x.get("event")] for x in dossier.get("timeline", [])])
    if dossier.get("human_impact"):
        _section(doc, "Human Impact", dossier.get("human_impact"))
    _section(doc, "Key Statistics")
    _table(doc, ["Metric", "Value", "Period", "Source"], [[x.get("metric"), x.get("value"), x.get("period"), x.get("source")] for x in dossier.get("key_statistics", [])])
    _section(doc, "Claims and Evidence")
    evidence_rows = []
    for x in dossier.get("claim_evidence", []):
        host = urlparse(x.get("source_url") or "").netloc or x.get("source_url") or "—"
        evidence_rows.append(["Verified" if x.get("verified") else "Unverified", x.get("claim"), x.get("note"), host])
    _table(doc, ["Status", "Claim", "Evidence note", "Source"], evidence_rows)
    _section(doc, "Open Questions")
    _bullets(doc, dossier.get("unresolved_questions"))
    _section(doc, "Risk and Sensitivity Review", (dossier.get("risk_review") or {}).get("explanation"))
    _bullets(doc, (dossier.get("risk_review") or {}).get("flags"))
    _section(doc, "Sources")
    source_rows = []
    for i, x in enumerate(dossier.get("sources", []), 1):
        host = urlparse(x.get("url") or "").netloc or x.get("url") or "—"
        source_rows.append([i, x.get("title"), x.get("publisher"), x.get("published"), f"{x.get('type', '')} · {x.get('reliability', '')}", host])
    _table(doc, ["#", "Source", "Publisher", "Date", "Type / reliability", "Link"], source_rows)
    if meta.get("report"):
        doc.add_section(WD_SECTION.NEW_PAGE)
        _add_original_report(doc, meta["report"])

    core = doc.core_properties
    core.title = str(topic.get("topic") or "Research Dossier")
    core.subject = "Validated editorial research dossier"
    core.author = str(site.get("name") or "Editorial Automation")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
