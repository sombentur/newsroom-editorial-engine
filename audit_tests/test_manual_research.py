"""Manual Editorial Workbench (owner request, 28 Sep 2026): one Topic ID -> one prompt -> one report link -> one
imported report. No database, browser or network: an in-memory collection and httpx.MockTransport stand in."""
import asyncio
import io
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import httpx
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

from lib import manual_research as m  # noqa: E402

LONG = "<article><h1>Report</h1>" + "<p>" + "Evidence from the ministry and field reporting. " * 80 + "</p></article>"


def _get(doc, key):
    for part in key.split("."):
        doc = doc.get(part) if isinstance(doc, dict) else None
    return doc


def _match(doc, flt):
    for key, cond in flt.items():
        value = _get(doc, key)
        if isinstance(cond, dict):
            if "$ne" in cond and value == cond["$ne"]:
                return False
            if "$in" in cond and value not in cond["$in"]:
                return False
            if "$nin" in cond and value in cond["$nin"]:
                return False
        elif value != cond:
            return False
    return True


def _set(doc, key, value):
    parts = key.split(".")
    for part in parts[:-1]:
        doc = doc.setdefault(part, {})
    doc[parts[-1]] = value


class Collection:
    def __init__(self, *docs):
        self.docs = [dict(d) for d in docs]

    async def find_one(self, flt, projection=None, sort=None):
        return next((dict(d) for d in self.docs if _match(d, flt)), None)

    async def update_one(self, flt, update):
        for d in self.docs:
            if _match(d, flt):
                for k, v in update.get("$set", {}).items():
                    _set(d, k, v)
                for k in update.get("$unset", {}):
                    d.pop(k, None)
                for k, v in update.get("$push", {}).items():
                    d.setdefault(k, []).append(v)
                return SimpleNamespace(modified_count=1)
        return SimpleNamespace(modified_count=0)

    async def update_many(self, flt, update):
        for d in [d for d in self.docs if _match(d, flt)]:
            for k, v in update.get("$set", {}).items():
                _set(d, k, v)

    async def insert_one(self, doc):
        self.docs.append(dict(doc))


def fake_db(*articles):
    return SimpleNamespace(articles=Collection(*articles), manual_imports=Collection(), browser_jobs=Collection())


def run(coro, db):
    with patch.object(m, "db", db), patch.object(m, "audit", AsyncMock()), patch.object(m, "ensure_runner"):
        return asyncio.run(coro)


def workbook(rows):
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "Research"
    sheet.append(m.COLUMNS)
    for topic_id, link in rows:
        sheet.append([topic_id, "t", "site", "English", "", "prompt", link, ""])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_the_downloaded_workbook_has_one_row_per_topic_and_reads_back():
    rows = [{"id": f"T{i}", "title": f"Topic {i}", "website": "English Edition", "language": "English",
             "category": "Economy", "prompt": f"Research prompt {i}", "link": "", "status_label": "Excel Exported"}
            for i in range(3)]
    content = m.build_workbook(rows)
    from openpyxl import load_workbook
    sheet = load_workbook(io.BytesIO(content))["Research"]
    assert [c.value for c in sheet[1]] == m.COLUMNS
    assert [sheet.cell(r, 6).value for r in range(2, 5)] == ["Research prompt 0", "Research prompt 1", "Research prompt 2"]
    parsed, errors = m.read_workbook(content)
    assert not errors and [p["topic_id"] for p in parsed] == ["T0", "T1", "T2"]


def test_workbook_checks_before_anything_is_imported():
    assert m.read_workbook(b"not a workbook")[1][0].startswith("This file is not an Excel workbook")
    from openpyxl import Workbook
    book = Workbook()
    book.active.append(["Topic", "Link"])
    out = io.BytesIO()
    book.save(out)
    errors = m.read_workbook(out.getvalue())[1]
    assert any("Topic ID" in e for e in errors) and any("Research Report Link" in e for e in errors)
    # A link inserted with display text: the link target is used.
    book = Workbook()
    sheet = book.active
    sheet.append(m.COLUMNS)
    sheet.append(["T1", "t", "", "", "", "", "report", ""])
    sheet.cell(2, 7).hyperlink = "https://gemini.google.com/app/026b0fa4de34145d"
    out = io.BytesIO()
    book.save(out)
    assert m.read_workbook(out.getvalue())[0][0]["link"] == "https://gemini.google.com/app/026b0fa4de34145d"


def test_links_are_classified():
    assert m.classify_link("https://gemini.google.com/app/026b0fa4de34145d?hl=en") == \
        ("https://gemini.google.com/app/026b0fa4de34145d", "gemini")
    assert m.classify_link("https://docs.google.com/document/d/abcdefghijk12/edit")[1] == "web"
    assert m.classify_link("https://g.co/gemini/share/abc")[0] is None, "share pages cannot be read"
    for bad in ("javascript:alert(1)", "report A", "ftp://x.example/r"):
        assert m.classify_link(bad)[0] is None


def test_each_row_maps_to_its_own_topic_and_errors_never_stop_the_others():
    """One Topic ID -> one link; unknown, duplicate or invalid rows are reported; existing research is never replaced."""
    articles = [
        {"id": "A", "stage": "selected", "site_key": "human", "topic_snapshot": {"topic": "Topic A"},
         "manual_research": {"status": "excel_exported", "prompt": "pa"}},
        {"id": "B", "stage": "selected", "site_key": "human", "topic_snapshot": {"topic": "Topic B"}},
        {"id": "C", "stage": "selected", "site_key": "human", "topic_snapshot": {"topic": "Topic C"},
         "dossier_meta": {"report": "old report"}, "manual_research": {"status": "research_imported", "link": "https://x.example/c"}},
        {"id": "D", "stage": "selected", "site_key": "human", "manual_research": {"status": "ready_to_import", "link": "https://x.example/d"}},
        {"id": "E", "stage": "selected", "site_key": "human"},
        {"id": "R", "stage": "rejected", "site_key": "human"},
    ]
    db = fake_db(*articles)
    content = workbook([
        ("A", "https://gemini.google.com/app/026b0fa4de34145d"),  # ready to import
        ("B", ""),  # no link yet: stays pending
        ("C", "https://x.example/c-new"),  # already has research: not replaced
        ("E", "https://x.example/d"),  # link belongs to D: skipped
        ("Z", "https://x.example/z"),  # unknown: not created
        ("E", "https://x.example/e"),  # E twice: both skipped
        ("R", "https://x.example/r"),  # rejected: skipped
        ("B2", "not a link"),
    ])
    db.articles.docs.append({"id": "B2", "stage": "selected", "site_key": "human"})
    summary = run(m.import_workbook(content, "research.xlsx", "owner"), db)
    results = {r["topic_id"] + str(r["row"]): r["result"] for r in summary["rows"]}
    by_id = {d["id"]: d for d in db.articles.docs}
    assert by_id["A"]["manual_research"]["status"] == "ready_to_import"
    assert by_id["A"]["manual_research"]["link"] == "https://gemini.google.com/app/026b0fa4de34145d"
    assert "manual_research" not in by_id["B"], "an empty link changes nothing"
    assert by_id["C"]["manual_research"]["status"] == "research_exists"
    assert by_id["C"]["manual_research"]["pending_link"] == "https://x.example/c-new"
    assert by_id["C"]["dossier_meta"]["report"] == "old report", "existing research is never silently replaced"
    assert "manual_research" not in by_id["E"], "duplicated Topic IDs and borrowed links are skipped"
    assert by_id["B2"]["manual_research"]["status"] == "link_invalid"
    assert results["Z6"].startswith("Unknown Topic ID") and not any(d["id"] == "Z" for d in db.articles.docs)
    assert results["E5"].startswith("Duplicate Topic ID") and results["E7"].startswith("Duplicate Topic ID")
    assert "rejected" in results["R8"]
    c = summary["counts"]
    assert c["total_rows"] == 8 and c["links_found"] == 7 and c["without_links"] == 1 and c["already_exists"] == 1
    assert c["in_progress"] == 1 and c["invalid_links"] == 1


def test_a_link_shared_by_two_rows_is_never_imported_twice():
    db = fake_db({"id": "A", "stage": "selected"}, {"id": "B", "stage": "selected"})
    summary = run(m.import_workbook(workbook([("A", "https://x.example/r"), ("B", "https://x.example/r")]), "f.xlsx", "o"), db)
    assert all("more than one row" in r["result"] for r in summary["rows"])
    assert not any(d.get("manual_research") for d in db.articles.docs)


def _web(handler):
    return httpx.AsyncClient(transport=httpx.MockTransport(handler), follow_redirects=True)


def test_report_links_are_read_and_problems_are_named():
    async def fetch(url, handler):
        async with _web(handler) as client:
            return await m._from_web(url, client)

    requested = []

    def doc(request):
        requested.append(str(request.url))
        return httpx.Response(200, headers={"content-type": "text/html"}, text="<html><body><nav>menu</nav>" + LONG + "</body></html>")
    content = asyncio.run(fetch("https://docs.google.com/document/d/abcdefghijk12/edit", doc))
    assert requested == ["https://docs.google.com/document/d/abcdefghijk12/export?format=html"], "a Google Doc's HTML export"
    assert "Evidence from the ministry" in content and "menu" not in content
    for status, expected in ((403, "access_denied"), (404, "report_not_found"), (500, "import_failed")):
        with pytest.raises(m.ImportProblem) as problem:
            asyncio.run(fetch("https://x.example/r", lambda r, s=status: httpx.Response(s)))
        assert problem.value.status == expected

    def login(request):
        if request.url.host == "accounts.google.com":
            return httpx.Response(200, headers={"content-type": "text/html"}, text="<p>Sign in</p>")
        return httpx.Response(302, headers={"location": "https://accounts.google.com/ServiceLogin?continue=x"})
    with pytest.raises(m.ImportProblem) as problem:
        asyncio.run(fetch("https://docs.google.com/document/d/abcdefghijk12/edit", login))
    assert problem.value.status == "access_denied", "a private document"


def test_a_topic_is_imported_only_when_report_content_was_retrieved():
    held = {"id": "A", "stage": "held_review", "held_reason": "research failed (browser): x", "site_key": "human",
            "manual_research": {"status": "ready_to_import", "link": "https://x.example/a", "prompt": "pa"}}
    db = fake_db(dict(held))
    with patch.object(m, "_from_web", AsyncMock(return_value="<p>too short</p>")):
        run(m.import_one(dict(held)), db)
    assert db.articles.docs[0]["manual_research"]["status"] == "report_not_found", "a URL alone is not an import"
    db = fake_db(dict(held))
    with patch.object(m, "_from_web", AsyncMock(return_value=LONG)):
        run(m.import_one(dict(held)), db)
    saved = db.articles.docs[0]
    assert saved["manual_research"]["status"] == "research_imported"
    assert saved["manual_research"]["link"] == "https://x.example/a", "the original link is kept"
    assert saved["dossier_meta"]["report"] == LONG and saved["dossier_meta"]["prompt"] == "pa"
    assert saved["stage"] == "selected" and saved["held_reason"] is None, "it continues in its turn"


def test_a_gemini_link_is_imported_through_the_extension_and_a_stopped_job_frees_the_lane():
    from lib import browser_bridge
    from lib.ai import AIError
    article = {"id": "A", "stage": "selected", "manual_research": {"status": "importing", "link": "u", "prompt": "pa"}}
    db = fake_db(dict(article))
    db.browser_config = SimpleNamespace(find_one=AsyncMock(return_value={"workspace": {"research": {"tab_id": 7}}}))
    with patch.object(browser_bridge, "_wait_for_job", AsyncMock(return_value=LONG)):
        assert run(m._from_gemini(article, "https://gemini.google.com/app/026b0fa4de34145d"), db) == LONG
    job = db.browser_jobs.docs[0]
    assert job["manual_import"] and job["capture_existing"] and job["capture_url"].endswith("026b0fa4de34145d")
    assert job["status"] == "consumed"
    db = fake_db(dict(article))
    db.browser_config = SimpleNamespace(find_one=AsyncMock(return_value={"workspace": {"research": {"tab_id": 7}}}))
    stuck = AIError("This Gemini conversation shows no finished report (no Share & Export).", "browser")
    with patch.object(browser_bridge, "_wait_for_job", AsyncMock(side_effect=stuck)), pytest.raises(m.ImportProblem) as problem:
        run(m._from_gemini(article, "https://gemini.google.com/app/026b0fa4de34145d"), db)
    assert problem.value.status == "report_not_found"


def test_reserved_topics_are_left_to_the_owner_and_release_the_turn():
    from lib import scheduler, turn
    reserved = {"id": "A", "stage": "held_review", "site_key": "human", "manual_research": {"status": "excel_exported"}}
    assert m.awaiting_manual(reserved)
    assert not m.awaiting_manual({**reserved, "dossier_meta": {"report": "r"}}), "imported: it continues"
    arts = SimpleNamespace(find=lambda *a, **k: SimpleNamespace(sort=lambda *a: SimpleNamespace(
        to_list=AsyncMock(return_value=[{**reserved, "stage": "selected"}, {"id": "B", "stage": "selected", "site_key": "human"}]))))
    with patch.object(scheduler, "db", SimpleNamespace(articles=arts)):
        article, step = asyncio.run(scheduler._next_work({"key": "human"}, {}))
    assert article["id"] == "B", "the automatic sequence skips a topic reserved for manual research"
    system = Collection({"id": "system", "current_article": "A"})
    with patch.object(turn, "db", SimpleNamespace(system_settings=system, articles=Collection(reserved))):
        assert asyncio.run(turn.current_article()) is None


def test_manual_imports_are_claimed_whichever_article_holds_the_turn():
    from lib import browser_bridge as b
    captured = {}

    async def claim(eligible, *args, **kwargs):
        captured["q"] = eligible
        return None
    request = SimpleNamespace(headers={"x-bridge-version": "0.4.28"}, json=AsyncMock(return_value={"lane": "research"}))
    database = SimpleNamespace(system_settings=SimpleNamespace(find_one=AsyncMock(return_value={"global_paused": False})),
                               browser_jobs=SimpleNamespace(find_one_and_update=claim))
    with patch.object(b, "db", database), patch("lib.turn.current_article", AsyncMock(return_value={"id": "CUR"})), \
            patch("lib.turn.next_article", AsyncMock(return_value=None)):
        asyncio.run(b.claim(request))
    manual = captured["q"]["$or"][1]
    assert manual == {"status": "queued", "manual_import": True, "kind": {"$in": ["research"]}}


def test_with_automatic_research_off_topics_wait_for_their_links():
    """Owner request (28 Sep 2026): the Editorial Workbench switch. Off, the app researches nothing by itself."""
    from lib import scheduler, turn
    waiting = {"id": "A", "stage": "selected", "site_key": "human"}
    imported = {"id": "B", "stage": "selected", "site_key": "human", "dossier_meta": {"report": "r"}}
    arts = SimpleNamespace(find=lambda *a, **k: SimpleNamespace(sort=lambda *a: SimpleNamespace(
        to_list=AsyncMock(return_value=[dict(waiting), dict(imported)]))))
    with patch.object(scheduler, "db", SimpleNamespace(articles=arts)):
        assert asyncio.run(scheduler._next_work({"key": "human"}, {"auto_research": False}))[0]["id"] == "B"
        assert asyncio.run(scheduler._next_work({"key": "human"}, {}))[0]["id"] == "A", "on: as before"
    held = {"id": "A", "stage": "held_review", "held_reason": "research failed (browser): x", "site_key": "human"}
    system = Collection({"id": "system", "current_article": "A", "auto_research": False})
    with patch.object(turn, "db", SimpleNamespace(system_settings=system, articles=Collection(held))):
        assert asyncio.run(turn.current_article()) is None, "a topic waiting for its link does not hold the queue"
    system = Collection({"id": "system", "current_article": "A", "auto_research": False})
    with patch.object(turn, "db", SimpleNamespace(system_settings=system, articles=Collection({**held, "stage": "researching"}))):
        assert asyncio.run(turn.current_article())["id"] == "A", "a research already running is not cut off"
    research = AsyncMock()
    fake = SimpleNamespace(articles=SimpleNamespace(find_one=AsyncMock(return_value=dict(waiting))),
                           system_settings=SimpleNamespace(find_one=AsyncMock(return_value={"auto_research": False})))
    with patch.object(scheduler, "db", fake), patch.object(scheduler, "run_research_stage", research), \
            patch.object(scheduler, "system_block_reason", AsyncMock(return_value=None)):
        asyncio.run(scheduler._complete_article(dict(waiting), {"key": "human"}))
    assert not research.called, "no research starts while automatic research is off"
    from models.schemas import SystemUpdate
    assert SystemUpdate(auto_research=False).model_dump(exclude_none=True) == {"auto_research": False}


def test_proceed_takes_a_pasted_link_and_puts_the_topic_first():
    """Owner request (28 Sep 2026): paste one topic's link in its row and Proceed, without the whole sheet."""
    db = fake_db({"id": "A", "stage": "selected", "site_key": "human", "manual_research": {"status": "excel_exported"}},
                 {"id": "B", "stage": "selected", "site_key": "human", "manual_research": {"status": "ready_to_import",
                                                                                          "link": "https://x.example/b"}},
                 {"id": "C", "stage": "selected", "site_key": "human", "dossier_meta": {"report": "r"},
                  "manual_research": {"status": "research_imported", "link": "https://x.example/c"}},
                 {"id": "D", "stage": "selected", "site_key": "human"})
    record = run(m.proceed("A", " https://gemini.google.com/app/026b0fa4de34145d?hl=en "), db)
    assert record["status"] == "ready_to_import" and record["link"] == "https://gemini.google.com/app/026b0fa4de34145d"
    assert record["proceed_at"], "it goes next once imported"
    with pytest.raises(ValueError, match="already belongs to another topic"):
        run(m.proceed("D", "https://x.example/b"), db)
    with pytest.raises(ValueError, match="Paste the research report link"):
        run(m.proceed("D"), db)
    with pytest.raises(ValueError, match="https://"):
        run(m.proceed("D", "not a link"), db)
    imported = run(m.proceed("C"), db)
    assert imported["status"] == "research_imported" and imported["proceed_at"], "research in: it just goes next"
    with pytest.raises(ValueError, match="already has research"):
        run(m.proceed("C", "https://x.example/other"), db)


def test_the_sequence_starts_the_proceed_topic_first_even_on_the_other_site():
    from datetime import datetime, timezone
    from lib import scheduler
    pressed = {"id": "P", "stage": "selected", "site_key": "kannadiga", "dossier_meta": {"report": "r"},
               "manual_research": {"status": "research_imported", "proceed_at": datetime.now(timezone.utc)}}
    older = {"id": "O", "stage": "selected", "site_key": "human", "dossier_meta": {"report": "r"}}
    by_site = {"kannadiga": [{"id": "K0", "stage": "selected", "site_key": "kannadiga", "dossier_meta": {"report": "r"}}, pressed],
               "human": [older]}

    def find(query, *a, **k):
        return SimpleNamespace(sort=lambda *a: SimpleNamespace(to_list=AsyncMock(return_value=[dict(x) for x in by_site[query["site_key"]]])))
    started = AsyncMock()
    with (patch.object(scheduler, "db", SimpleNamespace(articles=SimpleNamespace(find=find))),
          patch.object(scheduler, "is_connected", return_value=True),
          patch.object(scheduler, "current_article", AsyncMock(return_value=None)),
          patch.object(scheduler, "next_article", AsyncMock(return_value=None)),
          patch.object(scheduler, "_research_ahead", AsyncMock()),
          patch.object(scheduler, "_start", started),
          patch("routers.pipeline._research_tasks", {})):
        asyncio.run(scheduler._advance_sequence([{"key": "kannadiga"}, {"key": "human"}], {"last_sequence_site": "kannadiga"}))
    assert started.await_args.args[1]["id"] == "P", "Proceed goes first, although it is the other site's turn"
    by_site["kannadiga"][1] = {**pressed, "manual_research": {"status": "research_imported"}}
    started.reset_mock()
    with (patch.object(scheduler, "db", SimpleNamespace(articles=SimpleNamespace(find=find))),
          patch.object(scheduler, "is_connected", return_value=True),
          patch.object(scheduler, "current_article", AsyncMock(return_value=None)),
          patch.object(scheduler, "next_article", AsyncMock(return_value=None)),
          patch.object(scheduler, "_research_ahead", AsyncMock()),
          patch.object(scheduler, "_start", started),
          patch("routers.pipeline._research_tasks", {})):
        asyncio.run(scheduler._advance_sequence([{"key": "kannadiga"}, {"key": "human"}], {"last_sequence_site": "kannadiga"}))
    assert started.await_args.args[1]["id"] == "O", "without Proceed the sites alternate as before"


def test_the_owners_fetch_lifts_an_earlier_stop_and_a_recopy_replaces_research_at_once():
    """Live case (29 Sep 2026): after Stop every fetch was cancelled at once; a recopy threw away the imported report."""
    stopped = {"id": "A", "stage": "held_review", "held_reason": "Stopped by editor.", "site_key": "human",
               "manual_research": {"status": "import_failed", "link": "https://gemini.google.com/u/2/app/58b0469f5c8d26fe"}}
    db = fake_db(dict(stopped))
    run(m.proceed("A"), db)
    saved = db.articles.docs[0]
    assert saved["held_reason"] is None and saved["manual_research"]["status"] == "ready_to_import", "Stop is lifted"
    imported = {"id": "B", "stage": "article_validated", "site_key": "human", "held_reason": None,
                "dossier_meta": {"report_mode": True, "prompt": "p", "report": "old", "recopy_url": "x"},
                "article": {"h": 1}, "image": {"i": 1}}
    db = fake_db(dict(imported))
    db.browser_jobs.docs = [{"id": "J1", "article_id": "B", "kind": "seo", "status": "running"},
                            {"id": "J2", "article_id": "B", "kind": "research", "status": "queued", "manual_import": True}]
    record = run(m.import_again("B", "https://gemini.google.com/u/2/app/58b0469f5c8d26fe?hl=en"), db)
    saved = db.articles.docs[0]
    assert record["status"] == "ready_to_import" and record["proceed_at"], "fetched now; the topic goes next"
    assert record["link"] == "https://gemini.google.com/u/2/app/58b0469f5c8d26fe", "its account kept"
    assert saved["stage"] == "selected" and saved["article"] is None and saved["image"] is None
    assert saved["dossier_meta"]["report"] is None and "recopy_url" not in saved["dossier_meta"]
    assert db.browser_jobs.docs[0]["status"] == "cancelled" and db.browser_jobs.docs[1]["status"] == "queued"


def test_one_posts_research_row_is_the_same_as_its_manual_workbench_row():
    """Owner request (29 Sep 2026): the Editorial Workbench post shows its Manual Workbench row's box and actions."""
    queued = {"id": "Q", "stage": "selected", "site_key": "human", "topic_snapshot": {"topic": "Topic Q", "category": "Economy"}}
    db = fake_db(dict(queued))
    builder = AsyncMock(return_value=(lambda article: "Research prompt for Q", {"human": {"name": "English Edition", "language": "en"}}))
    with patch.object(m, "_prompt_builder", builder), patch.object(m, "_busy_topics", AsyncMock(return_value=set())):
        row = run(m.row_for("Q"), db)
        missing = run(m.row_for("nope"), db)
    assert row["prompt"] == "Research prompt for Q" and row["status"] == "pending" and row["exportable"]
    assert row["title"] == "Topic Q" and row["website"] == "English Edition" and missing is None
